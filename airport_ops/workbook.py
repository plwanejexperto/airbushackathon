"""Load and validate data/challenge2_operational_constraints.xlsx.

Sheets (actual names/columns, inspected from the file):
  airport_capacity : IATA_CODE, AIRPORT, CITY, STATE, delayed_departures_in_flights_csv,
                     size_class, active_runways, gates, max_runway_ops_per_hour,
                     min_turnaround_min, spare_gates_at_any_time   (SYNTHETIC limits)
  wake_separation  : leader, follower, min_separation_nm, min_time_s (same runway departures)
  fuel_and_crew    : rule, value (free text with units), applies_to
  rules_of_play    : topic, detail (free text)

Units are carried in column names / value text: *_min = minutes, *_nm = nautical miles,
*_s = seconds, "kg per hour", "hours". Preamble rows above headers are skipped by
locating the header row rather than assuming a row number.

Problems never raise: they are returned as WorkbookIssue entries.
"""
from __future__ import annotations

import re
from enum import Enum
from pathlib import Path
from typing import Any, Optional

import openpyxl
from pydantic import BaseModel, Field, ValidationError, model_validator

ROOT = Path(__file__).resolve().parent.parent
CONSTRAINTS_XLSX = ROOT / "data" / "challenge2_operational_constraints.xlsx"

SHEET_AIRPORTS = "airport_capacity"
SHEET_WAKE = "wake_separation"
SHEET_FUEL = "fuel_and_crew"
SHEET_RULES = "rules_of_play"
EXPECTED_SHEETS = (SHEET_AIRPORTS, SHEET_WAKE, SHEET_FUEL, SHEET_RULES)


# ----------------------------- models -----------------------------
class AircraftClass(str, Enum):
    NARROWBODY = "narrowbody"  # workbook: "A320/737 class"
    WIDEBODY = "widebody"      # workbook: "A330/777 class"


class WakeCategory(str, Enum):
    SUPER = "Super"
    HEAVY = "Heavy"
    LARGE = "Large"
    SMALL = "Small"


def parse_wake_category(text: str) -> WakeCategory:
    t = str(text).strip().lower()
    for cat in WakeCategory:
        if t.startswith(cat.value.lower()):  # "Super (A380)" -> SUPER
            return cat
    raise ValueError(f"unknown wake category {text!r}")


class Severity(str, Enum):
    ERROR = "error"
    WARNING = "warning"


class WorkbookIssue(BaseModel):
    severity: Severity
    sheet: str
    row: Optional[int] = None  # 1-based Excel row
    column: Optional[str] = None
    message: str

    def __str__(self) -> str:
        loc = self.sheet + (f" row {self.row}" if self.row else "") + (f" [{self.column}]" if self.column else "")
        return f"{self.severity.value.upper()} {loc}: {self.message}"


class AirportLimits(BaseModel):
    """Workbook-provided SYNTHETIC planning limits for one airport."""
    iata_code: str = Field(min_length=1)
    airport: str = Field(min_length=1)
    city: Optional[str] = None
    state: Optional[str] = None
    delayed_departures_in_flights_csv: int = Field(ge=0)
    size_class: str
    active_runways: int = Field(ge=1)
    gates: int = Field(ge=1)
    max_runway_ops_per_hour: int = Field(ge=1)  # movements/hour; see README for interpretation
    min_turnaround_min: int = Field(ge=1)       # minutes
    spare_gates_at_any_time: int = Field(ge=0)

    @model_validator(mode="after")
    def _spare_le_gates(self):
        if self.spare_gates_at_any_time > self.gates:
            raise ValueError("spare_gates_at_any_time exceeds gates")
        return self


class WakeRule(BaseModel):
    leader: WakeCategory
    follower: WakeCategory
    min_separation_nm: float = Field(gt=0)
    min_time_s: float = Field(gt=0)


class WakeTable(BaseModel):
    """Ordered (leader -> follower) rules plus the workbook's 'Any other pair' default."""
    rules: list[WakeRule]
    default_separation_nm: float = Field(gt=0)
    default_time_s: float = Field(gt=0)

    def lookup(self, leader: WakeCategory, follower: WakeCategory) -> tuple[float, float, str]:
        """Return (separation_nm, time_s, source label). Order matters: Large->Heavy is
        NOT the same as Heavy->Large."""
        for r in self.rules:
            if r.leader == leader and r.follower == follower:
                return r.min_separation_nm, r.min_time_s, f"{leader.value}->{follower.value}"
        return self.default_separation_nm, self.default_time_s, "Any other pair"


class FuelCrewRules(BaseModel):
    final_reserve_min: float = Field(gt=0)            # minutes of holding fuel
    contingency_fraction: float = Field(ge=0, le=1)   # fraction of trip fuel
    hold_burn_narrowbody_kg_h: float = Field(gt=0)
    hold_burn_widebody_kg_h: float = Field(gt=0)
    assumed_spare_fuel_min: float = Field(ge=0)       # minutes of holding fuel on arrival
    flight_crew_max_duty_h: float = Field(gt=0)       # hours from report time
    ground_shift_h: float = Field(gt=0)
    ground_overtime_multiplier: float = Field(ge=1)
    gate_change_notice_min: float = Field(ge=0)       # minutes before arrival

    def hold_burn_kg_h(self, cls: AircraftClass) -> float:
        return self.hold_burn_narrowbody_kg_h if cls == AircraftClass.NARROWBODY else self.hold_burn_widebody_kg_h


class RulesOfPlay(BaseModel):
    topics: dict[str, str]
    max_turns: Optional[int] = None


class WorkbookRules(BaseModel):
    source: str
    airports: dict[str, AirportLimits]
    invalid_airports: dict[str, str] = Field(default_factory=dict)  # code -> reason
    wake: WakeTable
    fuel_crew: FuelCrewRules
    rules_of_play: RulesOfPlay


class WorkbookLoadResult(BaseModel):
    rules: Optional[WorkbookRules] = None
    issues: list[WorkbookIssue] = Field(default_factory=list)

    @property
    def errors(self) -> list[WorkbookIssue]:
        return [i for i in self.issues if i.severity == Severity.ERROR]

    @property
    def usable(self) -> bool:
        return self.rules is not None


# ----------------------------- parsing helpers -----------------------------
def _num(v: Any) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _norm(h: Any) -> str:
    return str(h).split(" (")[0].strip() if h is not None else ""


def _find_header(rows, required: list[str]) -> Optional[tuple[int, dict[str, int]]]:
    for i, (_, cells) in enumerate(rows):
        cm = {_norm(c): j for j, c in enumerate(cells) if c is not None}
        if all(r in cm for r in required):
            return i, cm
    return None


def _sheet_rows(wb, name: str, issues: list[WorkbookIssue]):
    if name not in wb.sheetnames:
        issues.append(WorkbookIssue(severity=Severity.ERROR, sheet=name,
                                    message=f"required sheet missing (found: {wb.sheetnames})"))
        return None
    return [(i + 1, row) for i, row in enumerate(wb[name].iter_rows(values_only=True))]


def _err(issues, sheet, msg, row=None, col=None, sev=Severity.ERROR):
    issues.append(WorkbookIssue(severity=sev, sheet=sheet, row=row, column=col, message=msg))


# ----------------------------- sheet loaders -----------------------------
_AIRPORT_COLS = ["IATA_CODE", "AIRPORT", "CITY", "STATE", "delayed_departures_in_flights_csv", "size_class",
                 "active_runways", "gates", "max_runway_ops_per_hour", "min_turnaround_min",
                 "spare_gates_at_any_time"]


def _load_airports(wb, issues):
    rows = _sheet_rows(wb, SHEET_AIRPORTS, issues)
    if rows is None:
        return None, {}
    hdr = _find_header(rows, _AIRPORT_COLS)
    if hdr is None:
        present = {_norm(c) for _, r in rows for c in r if c}
        _err(issues, SHEET_AIRPORTS, f"header row not found; missing columns: "
             f"{[c for c in _AIRPORT_COLS if c not in present]}")
        return None, {}
    hi, cm = hdr
    out: dict[str, AirportLimits] = {}
    bad: dict[str, str] = {}
    for rown, cells in rows[hi + 1:]:
        if all(c is None for c in cells):
            continue
        rec = {c: cells[cm[c]] for c in _AIRPORT_COLS}
        code = str(rec["IATA_CODE"]).strip() if rec["IATA_CODE"] is not None else ""
        try:
            lim = AirportLimits(
                iata_code=code, airport=str(rec["AIRPORT"] or "").strip(),
                city=rec["CITY"], state=rec["STATE"],
                delayed_departures_in_flights_csv=rec["delayed_departures_in_flights_csv"],
                size_class=str(rec["size_class"] or ""), active_runways=rec["active_runways"],
                gates=rec["gates"], max_runway_ops_per_hour=rec["max_runway_ops_per_hour"],
                min_turnaround_min=rec["min_turnaround_min"], spare_gates_at_any_time=rec["spare_gates_at_any_time"])
        except ValidationError as e:
            msg = "; ".join(f"{'.'.join(map(str, x['loc']))}: {x['msg']}" for x in e.errors())
            _err(issues, SHEET_AIRPORTS, f"invalid row for {code or '<blank code>'}: {msg}", rown)
            if code:
                bad[code] = msg
            continue
        if code in out:
            _err(issues, SHEET_AIRPORTS, f"duplicate IATA_CODE {code}; first occurrence kept", rown, "IATA_CODE")
            continue
        out[code] = lim
    if not out:
        _err(issues, SHEET_AIRPORTS, "no valid airport rows")
        return None, bad
    return out, bad


def _load_wake(wb, issues) -> Optional[WakeTable]:
    rows = _sheet_rows(wb, SHEET_WAKE, issues)
    if rows is None:
        return None
    hdr = _find_header(rows, ["leader", "follower", "min_separation_nm", "min_time_s"])
    if hdr is None:
        _err(issues, SHEET_WAKE, "header row not found (need leader, follower, min_separation_nm, min_time_s)")
        return None
    hi, cm = hdr
    rules, seen, default = [], set(), None
    n_err = len([i for i in issues if i.severity == Severity.ERROR])
    for rown, cells in rows[hi + 1:]:
        if all(c is None for c in cells):
            continue
        leader, follower = cells[cm["leader"]], cells[cm["follower"]]
        nm, ts = _num(cells[cm["min_separation_nm"]]), _num(cells[cm["min_time_s"]])
        if nm is None or nm <= 0 or ts is None or ts <= 0:
            _err(issues, SHEET_WAKE, f"separation/time must be positive numbers (got {cells[cm['min_separation_nm']]!r}, "
                 f"{cells[cm['min_time_s']]!r})", rown)
            continue
        if leader is not None and str(leader).strip().lower().startswith("any other"):
            if default is not None:
                _err(issues, SHEET_WAKE, "duplicate 'Any other pair' row", rown)
            default = (nm, ts)
            continue
        try:
            l, f = parse_wake_category(leader), parse_wake_category(follower)
        except ValueError as e:
            _err(issues, SHEET_WAKE, str(e), rown)
            continue
        if (l, f) in seen:
            _err(issues, SHEET_WAKE, f"duplicate pair {l.value}->{f.value}", rown)
            continue
        seen.add((l, f))
        rules.append(WakeRule(leader=l, follower=f, min_separation_nm=nm, min_time_s=ts))
    if default is None:
        _err(issues, SHEET_WAKE, "missing 'Any other pair' default row")
    if default is None or len([i for i in issues if i.severity == Severity.ERROR]) > n_err:
        return None
    return WakeTable(rules=rules, default_separation_nm=default[0], default_time_s=default[1])


_N = r"(\d+(?:\.\d+)?)"
# rule name -> (regex over value text, target fields (group order))
_FUEL_RULES: dict[str, tuple[str, tuple[str, ...]]] = {
    "Final fuel reserve": (rf"^\s*{_N}\s*minutes?\s+of\s+holding\s+fuel", ("final_reserve_min",)),
    "Contingency fuel": (rf"^\s*{_N}\s*%\s*of\s+trip\s+fuel", ("contingency_fraction",)),
    "Holding fuel burn (narrowbody)": (r"^\s*([\d,]+(?:\.\d+)?)\s*kg\s+per\s+hour", ("hold_burn_narrowbody_kg_h",)),
    "Holding fuel burn (widebody)": (r"^\s*([\d,]+(?:\.\d+)?)\s*kg\s+per\s+hour", ("hold_burn_widebody_kg_h",)),
    "Assumed spare fuel on arrival": (rf"^\s*{_N}\s*minutes?\s+of\s+holding", ("assumed_spare_fuel_min",)),
    "Flight crew max duty": (rf"^\s*{_N}\s*hours?\s+from\s+report\s+time", ("flight_crew_max_duty_h",)),
    "Ground crew shift": (rf"^\s*{_N}\s*hours?\s*;\s*overtime\s+costs\s+{_N}\s*x\s+after",
                          ("ground_shift_h", "ground_overtime_multiplier")),
    "Gate change notice": (rf"^\s*at\s+least\s+{_N}\s*minutes?\s+before\s+arrival", ("gate_change_notice_min",)),
}


def _load_fuel_crew(wb, issues) -> Optional[FuelCrewRules]:
    rows = _sheet_rows(wb, SHEET_FUEL, issues)
    if rows is None:
        return None
    hdr = _find_header(rows, ["rule", "value", "applies_to"])
    if hdr is None:
        _err(issues, SHEET_FUEL, "header row not found (need rule, value, applies_to)")
        return None
    hi, cm = hdr
    text = {str(c[cm["rule"]]).strip(): (r, c[cm["value"]]) for r, c in rows[hi + 1:] if c[cm["rule"]] is not None}
    vals: dict[str, float] = {}
    ok = True
    for name, (rx, fields) in _FUEL_RULES.items():
        if name not in text:
            _err(issues, SHEET_FUEL, f"required rule {name!r} missing", col="rule")
            ok = False
            continue
        rown, raw = text[name]
        m = re.search(rx, str(raw), re.I)
        if not m:
            _err(issues, SHEET_FUEL, f"cannot parse {name!r} value {raw!r} (expected pattern {rx})", rown, "value")
            ok = False
            continue
        for f, g in zip(fields, m.groups()):
            vals[f] = float(g.replace(",", ""))
    if not ok:
        return None
    vals["contingency_fraction"] /= 100.0  # "5% of trip fuel"
    try:
        return FuelCrewRules(**vals)
    except ValidationError as e:
        _err(issues, SHEET_FUEL, "; ".join(f"{x['loc']}: {x['msg']}" for x in e.errors()))
        return None


def _load_rules_of_play(wb, issues) -> Optional[RulesOfPlay]:
    rows = _sheet_rows(wb, SHEET_RULES, issues)
    if rows is None:
        return None
    hdr = _find_header(rows, ["topic", "detail"])
    if hdr is None:
        _err(issues, SHEET_RULES, "header row not found (need topic, detail)")
        return None
    hi, cm = hdr
    topics = {str(c[cm["topic"]]).strip(): str(c[cm["detail"]]).strip()
              for _, c in rows[hi + 1:] if c[cm["topic"]] is not None}
    m = re.search(r"within\s+(\d+)\s+turns?", topics.get("Turn", ""), re.I)
    if not m:
        _err(issues, SHEET_RULES, "could not find 'within N turns' in the Turn rule", col="detail", sev=Severity.WARNING)
    return RulesOfPlay(topics=topics, max_turns=int(m.group(1)) if m else None)


def load_workbook_rules(path: Path | str = CONSTRAINTS_XLSX) -> WorkbookLoadResult:
    issues: list[WorkbookIssue] = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # missing / corrupt file
        _err(issues, str(path), f"cannot open workbook: {e}")
        return WorkbookLoadResult(issues=issues)
    for extra in [s for s in wb.sheetnames if s not in EXPECTED_SHEETS]:
        _err(issues, extra, "unrecognised sheet ignored", sev=Severity.WARNING)
    airports, bad = _load_airports(wb, issues)
    wake = _load_wake(wb, issues)
    fuel = _load_fuel_crew(wb, issues)
    rop = _load_rules_of_play(wb, issues)
    wb.close()
    if airports is None or wake is None or fuel is None or rop is None:
        return WorkbookLoadResult(issues=issues)
    return WorkbookLoadResult(
        rules=WorkbookRules(source=str(path), airports=airports, invalid_airports=bad,
                            wake=wake, fuel_crew=fuel, rules_of_play=rop),
        issues=issues)
