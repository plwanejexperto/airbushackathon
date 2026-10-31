import shutil
from pathlib import Path

import openpyxl
import pytest

from airport_ops import workbook as wbk
from airport_ops.scenario import ScenarioError, build_scenario
from airport_ops.workbook import WakeCategory as W, load_workbook_rules


def test_real_workbook_loads_all_sheets(rules):
    assert len(rules.airports) == 322
    atl = rules.airports["ATL"]
    assert (atl.active_runways, atl.gates, atl.max_runway_ops_per_hour, atl.min_turnaround_min,
            atl.spare_gates_at_any_time) == (4, 150, 90, 45, 3)
    f = rules.fuel_crew
    assert (f.final_reserve_min, f.contingency_fraction, f.hold_burn_narrowbody_kg_h, f.hold_burn_widebody_kg_h,
            f.assumed_spare_fuel_min, f.flight_crew_max_duty_h, f.ground_shift_h, f.ground_overtime_multiplier,
            f.gate_change_notice_min) == (30, 0.05, 2400, 6500, 45, 13, 8, 1.5, 20)
    assert rules.rules_of_play.max_turns == 3


def test_wake_lookup_is_ordered(rules):
    wk = rules.wake
    assert wk.lookup(W.HEAVY, W.LARGE)[:2] == (5, 120)
    assert wk.lookup(W.SUPER, W.SMALL)[:2] == (8, 180)
    assert wk.lookup(W.LARGE, W.HEAVY)[:3] == (3, 60, "Any other pair")  # reverse order is NOT the Heavy->Large rule
    assert wk.lookup(W.LARGE, W.LARGE)[2] == "Any other pair"


def _copy_and_edit(tmp_path, edit) -> Path:
    dst = tmp_path / "w.xlsx"
    shutil.copy(wbk.CONSTRAINTS_XLSX, dst)
    book = openpyxl.load_workbook(dst)
    edit(book)
    book.save(dst)
    return dst


def test_invalid_airport_row_is_reported_and_blocks_only_that_airport(tmp_path, rules, flights_df):
    def edit(b):
        ws = b["airport_capacity"]
        for row in ws.iter_rows(min_row=4):
            if row[0].value == "ATL":
                row[9].value = "abc"  # min_turnaround_min
    res = load_workbook_rules(_copy_and_edit(tmp_path, edit))
    assert res.usable and "ATL" in res.rules.invalid_airports
    assert any(i.sheet == "airport_capacity" and "ATL" in i.message and i.row for i in res.errors)
    from airport_ops import data
    h = data.to_historical(flights_df.loc[data.default_index(flights_df)])
    with pytest.raises(ScenarioError, match="ATL"):
        build_scenario(h, res.rules)


def test_missing_wake_default_and_bad_fuel_text_are_fatal_and_clear(tmp_path):
    def edit(b):
        ws = b["wake_separation"]
        for row in ws.iter_rows(min_row=4):
            if row[0].value == "Any other pair":
                row[0].value = "Large"; row[1].value = "Large"
        for row in b["fuel_and_crew"].iter_rows(min_row=2):
            if row[0].value == "Gate change notice":
                row[1].value = "whenever"
    res = load_workbook_rules(_copy_and_edit(tmp_path, edit))
    assert not res.usable
    msgs = " | ".join(str(i) for i in res.errors)
    assert "Any other pair" in msgs and "Gate change notice" in msgs and "cannot parse" in msgs


def test_missing_sheet_and_missing_file(tmp_path):
    res = load_workbook_rules(_copy_and_edit(tmp_path, lambda b: b.remove(b["fuel_and_crew"])))
    assert not res.usable and any("fuel_and_crew" in str(i) and "missing" in str(i) for i in res.errors)
    res = load_workbook_rules(tmp_path / "nope.xlsx")
    assert not res.usable and "cannot open" in str(res.errors[0])
